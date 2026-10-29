"""
Project knowledge base: examples (source documents + the expected summary) and the summary profile
learned from them. Everything lives in the project folder so it travels with it:

    <folder>/.workbench/projects/<id>/knowledge/  (.workbench/knowledge/ for older single-project folders)
        examples/<id>/example.json   {id, name, created, origin: upload|chat, sources: [file names], summary_file}
        examples/<id>/<source files>, summary.md   (the expected summary as Markdown)
        examples/<id>/evidence.json  alignment evidence measured in code (alignment.py), cached
        examples/<id>/notes.json     the LLM's analysis of that example, cached (re-learning only analyses new ones)
        profile.json                 {measured, targets, pairs, learned, pinned, text, version, history, ...}

Learning, in three steps:
  1. code aligns each summary sentence with its source passage and measures length, structure, coverage,
     verbatim share, numbers and bolding (evidence);
  2. the LLM analyses each example from the aligned pairs (why things are kept, how they are rewritten, tone);
  3. code aggregates the measurements across examples (with support counts) and the LLM merges the analyses
     into qualitative rules, each with its support (k/N).
The user reviews the draft; rules they pin are kept when re-learning. A confirmed profile, plus a few
"source -> summary" rewrite examples, goes into every chat in the project, and drafts are checked against
its measured targets (check_draft).
"""
import json
import re
import shutil
import time
import uuid
from pathlib import Path

import pymupdf

from . import alignment, citations, documents, llm, projects
from . import files as fs

SUMMARY_TYPES = {".md", ".txt", ".docx", ".pdf"}

EXAMPLE_PROMPT = """You are studying how an analyst turns source documents into a summary, so that the same kind of summary can be written for new documents. The app has already aligned each summary sentence with the source passage it comes from and measured the facts below; do not recount them.

Measured: {metrics}
Source sections used: {used}
Source sections skipped: {skipped}

Aligned sentences (source passage -> summary sentence):
{pairs}

Summary sentences with no clear source (the analyst's own additions):
{unmatched}

The full summary:
<summary>
{summary}
</summary>

Describe concisely, with short quotes as evidence:
1. Content selection: which kinds of information are kept and which are skipped, as reusable rules, and why.
2. Rewriting: how source sentences become summary sentences (what is shortened, merged, reordered, kept verbatim, abbreviated).
3. Tone and phrasing: tense, voice, formality, typical openings, abbreviations, number and date formats.
4. What the analyst adds that is not in the source (judgements, context), if anything.
5. Facts this kind of summary always needs.
Do not describe citation or reference formatting."""

MERGE_PROMPT = """Below are analyses of {n} example(s) of the same summary task.

These facts were measured exactly by the app and are already part of the profile; do not restate the numbers:
{measured}

Current project instructions:
<instructions>
{instructions}
</instructions>

<analyses>
{analyses}
</analyses>

Write the qualitative rules a writer needs to produce indistinguishable summaries for new documents. End every rule with its support in parentheses: (3/3) when every example shows it, (1/3) when only one does. Keep rules supported by at least half of the examples, plus any rule that is clearly important. Quote a short example phrase for each style rule. Do not write rules about citations, reference numbers, footnotes or source lists: the app adds citations itself. Output exactly this Markdown and nothing else:

## Learned rules
### Required information
(facts every summary needs; the assistant asks the user when one is missing from the documents)
### Content selection
### Rewriting
### Tone and phrasing
### Formatting
### Template
(a skeleton of the summary with [placeholders])

## Suggested additions to the project instructions
(3-6 bullets the user may add to the instructions; skip anything the current instructions already say)"""


# ---------------------------------------------------------------- storage

def _kdir(project: dict) -> Path:
    d = projects.meta_dir(project["folder"], project["id"]) / "knowledge"
    (d / "examples").mkdir(parents=True, exist_ok=True)
    return d


def _safe(name: str) -> str:
    return re.sub(r'[^\w.\- ]', "_", Path(name).name).strip() or "file"


def list_examples(pid: str) -> list[dict]:
    d = _kdir(projects.get(pid)) / "examples"
    out = []
    for f in sorted(d.glob("*/example.json")):
        ex = json.loads(f.read_text("utf-8"))
        ex["summary"] = (f.parent / ex["summary_file"]).read_text("utf-8") if (f.parent / ex["summary_file"]).exists() else ""
        out.append(ex)
    return sorted(out, key=lambda e: e["created"])


def add_example(pid: str, name: str, sources: list[tuple[str, bytes]], summary_md: str,
                origin: str = "upload") -> dict:
    """Store an example: source files (copied) + the expected summary (Markdown text)."""
    if not summary_md.strip():
        raise ValueError("The expected summary is empty.")
    for fname, _ in sources:
        if Path(fname).suffix.lower() not in fs.SUPPORTED:
            raise ValueError(f"{fname}: source documents must be PDF, Word, Text or Markdown files.")
    eid = time.strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:4]
    d = _kdir(projects.get(pid)) / "examples" / eid
    d.mkdir(parents=True)
    names = []
    for fname, data in sources:
        target = d / _safe(fname)
        if target.name == "summary.md" or target.exists():
            target = d / f"source-{len(names) + 1}-{_safe(fname)}"
        target.write_bytes(data)
        names.append(target.name)
    (d / "summary.md").write_text(summary_md.strip() + "\n", "utf-8")
    ex = {"id": eid, "name": name.strip() or (Path(sources[0][0]).stem if sources else f"Example {eid}"),
          "created": time.time(), "origin": origin, "sources": names, "summary_file": "summary.md"}
    (d / "example.json").write_text(json.dumps(ex, indent=1, ensure_ascii=False), "utf-8")
    return {**ex, "summary": summary_md}


def add_example_from_chat(pid: str, chat: dict, summary_md: str, name: str) -> dict:
    """An answer the user approved. Its sources: the documents attached to the chat, plus project context
    documents the summary cites (e.g. a memo added as C1)."""
    cited = {m.group(1).upper() for m in re.finditer(r"\b([CD]\d+)-[a-z]?\d+", summary_md, re.I)}
    entries = list(chat.get("docs", []))
    entries += [c for c in projects.get(pid)["context"] if c["label"].upper() in cited]
    sources = []
    for d in entries:
        try:
            src = documents.source_file(documents.get(d["doc_id"]))
            sources.append((d["name"], src.read_bytes()))
        except (FileNotFoundError, OSError):
            continue
    return add_example(pid, name, sources, summary_md, origin="chat")


def remove_example(pid: str, eid: str):
    d = _kdir(projects.get(pid)) / "examples" / eid
    if d.parent.name == "examples" and (d / "example.json").exists():
        shutil.rmtree(d)


def get_profile(project: dict) -> dict | None:
    f = _kdir(project) / "profile.json"
    if not f.exists():
        return None
    prof = json.loads(f.read_text("utf-8"))
    if "learned" not in prof:                       # profiles saved before measured/pinned existed
        prof.update({"learned": prof.get("text", ""), "pinned": "", "measured_md": "", "targets": {}, "pairs": []})
    return prof


def render_profile(measured_md: str, learned: str, pinned: str) -> str:
    parts = []
    if measured_md.strip():
        parts.append(measured_md.strip())
    if learned.strip():
        parts.append(learned.strip())
    if pinned.strip():
        parts.append("## Pinned rules (always apply; these override the rules above)\n" + pinned.strip())
    return "\n\n".join(parts)


def _changes(old: dict | None, new: dict) -> str:
    """One line describing what a new version changed."""
    if not old:
        return "first version"
    out = []
    ot, nt = old.get("targets") or {}, new.get("targets") or {}
    if list(ot.get("words") or []) != list(nt.get("words") or []) and nt.get("words"):
        out.append(f"length target {_rng(ot.get('words'))} → {_rng(nt.get('words'))} words")
    if ot.get("examples") != nt.get("examples") and nt.get("examples"):
        out.append(f"examples {ot.get('examples', 0)} → {nt['examples']}")
    rule = lambda t: {ln.strip() for ln in t.splitlines() if ln.strip().startswith(("-", "*"))}
    added = len(rule(new["learned"]) - rule(old.get("learned", "")))
    removed = len(rule(old.get("learned", "")) - rule(new["learned"]))
    if added or removed:
        out.append(f"learned rules +{added} / −{removed}")
    if old.get("pinned", "").strip() != new["pinned"].strip():
        out.append("pinned rules edited")
    return "; ".join(out) or "no rule changes"


def save_profile(pid: str, learned: str, pinned: str = "", draft: dict | None = None, source: str = "learned") -> dict:
    """Save a confirmed profile as a new version (the previous one goes to history). `draft` carries the
    measured part from `learn`; without it, the measured part of the current profile is kept."""
    project = projects.get(pid)
    old = get_profile(project)
    base = draft or old or {}
    prof = {"learned": learned.strip(), "pinned": pinned.strip(),
            "measured_md": base.get("measured_md", ""), "targets": base.get("targets", {}),
            "pairs": base.get("pairs", []),
            "version": (old or {}).get("version", 0) + 1, "updated": time.time(), "source": source}
    prof["text"] = render_profile(prof["measured_md"], prof["learned"], prof["pinned"])
    prof["changes"] = _changes(old, prof)
    history = (old or {}).get("history", [])
    if old:
        history.append({k: old.get(k) for k in ("text", "learned", "pinned", "version", "updated", "source", "changes")})
    prof["history"] = history[-20:]
    (_kdir(project) / "profile.json").write_text(json.dumps(prof, indent=1, ensure_ascii=False), "utf-8")
    return prof


def clear_profile(pid: str):
    (_kdir(projects.get(pid)) / "profile.json").unlink(missing_ok=True)


# ---------------------------------------------------------------- reading summaries

def summary_markdown(name: str, data: bytes) -> str:
    """Text of an expected-summary file as Markdown, keeping headings, bullets and bold (Word)."""
    suffix = Path(name).suffix.lower()
    if suffix in (".md", ".txt"):
        return data.decode("utf-8", "replace")
    if suffix == ".docx":
        return _docx_markdown(data)
    if suffix == ".pdf":
        pdf = pymupdf.open(stream=data, filetype="pdf")
        return "\n\n".join(page.get_text().strip() for page in pdf)
    raise ValueError(f"{name}: the expected summary must be a Word, PDF, Text or Markdown file.")


def _docx_markdown(data: bytes) -> str:
    import io
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        text = "".join(f"**{r.text}**" if r.bold and r.text.strip() else r.text for r in p.runs).strip()
        text = text.replace("****", "")
        if not text:
            continue
        style = (p.style.name or "").lower()
        m = re.match(r"heading\s*(\d)", style)
        numbered = p._p.pPr is not None and p._p.pPr.numPr is not None
        if m or style == "title":
            lines.append(f"{'#' * (int(m.group(1)) if m else 1)} {text.strip('*')}")
        elif "list" in style or numbered:
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    for t in d.tables:
        rows = [[c.text.strip().replace("|", "/") for c in r.cells] for r in t.rows]
        if rows:
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + "---|" * len(rows[0]))
            lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            lines.append("")
    return "\n".join(lines).strip()


# ---------------------------------------------------------------- evidence (step 1, code)

def _example_docs(kdir: Path, ex: dict) -> list[tuple[str, dict]]:
    docs = []
    for i, name in enumerate(ex["sources"], 1):
        try:
            docs.append((f"S{i}", documents.ingest(kdir / ex["id"] / name, name=name)))
        except Exception:
            continue
    return docs


def _key(ex: dict, extra: str = "") -> str:
    import hashlib
    return hashlib.sha1(json.dumps([ex["summary"], ex["sources"], alignment.ALIGN_VERSION, extra]).encode()).hexdigest()[:16]


def evidence(pid: str, ex: dict) -> dict:
    """Alignment evidence for one example (cached in the example folder)."""
    kdir = _kdir(projects.get(pid)) / "examples"
    f = kdir / ex["id"] / "evidence.json"
    key = _key(ex)
    if f.exists():
        cached = json.loads(f.read_text("utf-8"))
        if cached.get("key") == key:
            return cached["evidence"]
    ev = alignment.align(ex["summary"], _example_docs(kdir, ex))
    f.write_text(json.dumps({"key": key, "evidence": ev}, ensure_ascii=False), "utf-8")
    return ev


# ---------------------------------------------------------------- aggregation (code)

def _rng(r) -> str:
    if not r:
        return "?"
    lo, hi = r
    return f"{lo}" if lo == hi else f"{lo}–{hi}"


def _span(values, fmt=lambda v: v):
    return (fmt(min(values)), fmt(max(values))) if values else None


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower())


def aggregate(evs: list[dict]) -> dict:
    """Measured facts across examples, with support counts, plus targets for checking drafts."""
    n = len(evs)
    order, sections = [], {}
    for ev in evs:
        for sec in ev["sections"]:
            k = _norm(sec["name"])
            if k not in sections:
                order.append(k)
                sections[k] = {"name": sec["name"] or "(no heading)", "examples": 0, "points": []}
            sections[k]["points"].append(sec["points"])
        for k in {_norm(sec["name"]) for sec in ev["sections"]}:
            sections[k]["examples"] += 1
    src_map: dict[str, dict[str, int]] = {}
    for ev in evs:
        for sec, srcs in ev["source_map"].items():
            for src in srcs:
                m = src_map.setdefault(_norm(sec), {})
                m[src] = m.get(src, 0) + 1
    coverage: dict[str, dict] = {}
    for ev in evs:
        for src, c in ev["coverage"].items():
            if c["words"] < 15:                      # titles, sign-off lines, tiny fragments
                continue
            row = coverage.setdefault(src, {"present": 0, "used": 0})
            row["present"] += 1
            row["used"] += 1 if c["used"] else 0
    pct = lambda v: round(v * 100)
    words = [ev["words"] for ev in evs]
    verbatim = [ev["verbatim"] for ev in evs]
    nums = sum(ev["numbers"] for ev in evs)
    missing = sum(ev["numbers_missing"] for ev in evs)
    kinds: dict[str, int] = {}
    for ev in evs:
        for k, v in ev["bold"]["kinds"].items():
            kinds[k] = kinds.get(k, 0) + v
    pairs = [p for ev in reversed(evs) for p in ev["pairs"]]
    chosen, seen = [], set()
    for p in pairs + pairs:                          # one pair per summary section first, then fill up
        if len(chosen) >= 4:
            break
        if p in chosen or (p["section"] in seen and len(seen) < len({q["section"] for q in pairs})):
            continue
        seen.add(p["section"])
        chosen.append(p)
    return {
        "examples": n,
        "sections": [sections[k] for k in order],
        "source_map": src_map,
        "coverage": coverage,
        "words": _span(words), "words_per_point": _span([ev["words_per_point"] for ev in evs]),
        "points": _span([ev["points"] for ev in evs]),
        "compression": _span([ev["compression"] for ev in evs], pct),
        "verbatim": _span(verbatim, pct),
        "numbers_found": pct(1 - missing / nums) if nums else None,
        "unsourced": round(sum(len(ev["unmatched"]) for ev in evs) / n, 1) if n else 0,
        "bold": _span([ev["bold"]["count"] for ev in evs]), "bold_kinds": kinds,
        "pairs": chosen,
        "targets": {
            "examples": n,
            "words": (int(min(words) * 0.85), int(max(words) * 1.15 + 0.5)) if words else None,
            "verbatim": (max(0, pct(min(verbatim)) - 10), min(100, pct(max(verbatim)) + 10)) if verbatim else None,
            "sections": [sections[k]["name"] for k in order if sections[k]["examples"] * 2 >= n and sections[k]["name"]],
            "required_sources": [src for src, c in coverage.items()
                                 if c["used"] == c["present"] and (c["present"] >= 2 or n == 1)],
            "unsourced": round(sum(len(ev["unmatched"]) for ev in evs) / n, 1) if n else 0,
        },
    }


def measured_markdown(m: dict) -> str:
    n = m["examples"]
    lines = [f"## Measured from {n} example{'s' if n != 1 else ''} (by the app)", "### Structure"]
    if m["sections"]:
        lines.append("- Sections in order: " + "; ".join(
            f"**{s['name']}** — {_rng(_span(s['points']))} points ({s['examples']}/{n})" for s in m["sections"]))
    for s in m["sections"]:
        srcs = m["source_map"].get(_norm(s["name"]), {})
        if srcs:
            top = sorted(srcs.items(), key=lambda kv: -kv[1])[:6]
            lines.append(f"- **{s['name']}** draws on: " + ", ".join(f"{k} ({v}/{n})" for k, v in top))
    always = [f"{k} ({c['used']}/{c['present']})" for k, c in m["coverage"].items() if c["used"] == c["present"]]
    some = [f"{k} ({c['used']}/{c['present']})" for k, c in m["coverage"].items() if 0 < c["used"] < c["present"]]
    never = [f"{k} (0/{c['present']})" for k, c in m["coverage"].items() if c["used"] == 0]
    lines.append("### Content coverage (source sections)")
    if always:
        lines.append("- Always used: " + ", ".join(always))
    if some:
        lines.append("- Sometimes used: " + ", ".join(some))
    if never:
        lines.append("- Never used: " + ", ".join(never))
    lines.append(f"- About {m['unsourced']:g} sentence(s) per summary have no clear source (the writer's own context)"
                 if m["unsourced"] else "- Every sentence can be traced to the source")
    lines += ["### Length",
              f"- {_rng(m['words'])} words per summary (about {_rng(m['compression'])}% of the source's length); "
              f"{_rng(m['points'])} points; about {_rng(m['words_per_point'])} words per point",
              "### Rewriting",
              f"- About {_rng(m['verbatim'])}% of the wording is copied verbatim from the source (runs of 4+ words); "
              "the rest is rewritten"]
    if m["numbers_found"] is not None:
        lines.append(f"- {m['numbers_found']}% of the figures appear in the source (exactly or rounded)")
    if m["bold"]:
        kinds = ", ".join(f"{k} ({v})" for k, v in sorted(m["bold_kinds"].items(), key=lambda kv: -kv[1]))
        lines += ["### Bolding", f"- {_rng(m['bold'])} bold phrases per summary" + (f": {kinds}" if kinds else "")]
    return "\n".join(lines)


# ---------------------------------------------------------------- learning (steps 2 and 3, LLM)

def _example_prompt(ex: dict, ev: dict) -> str:
    used = [f"{k} ({c['used']}/{c['total']} paragraphs)" for k, c in ev["coverage"].items() if c["used"]]
    skipped = [k for k, c in ev["coverage"].items() if not c["used"] and c["words"] >= 15]
    pairs = "\n".join(
        f"[{r['section'] or 'summary'} ← {r['source_section']}, p.{r['page']}, {round(r['verbatim'] * 100)}% verbatim]\n"
        f"  SOURCE: {r['source_sentence'][:350]}\n  SUMMARY: {r['text']}" for r in ev["sentences"])
    metrics = (f"{ev['words']} words, {ev['points']} points in {len(ev['sections'])} section(s), "
               f"{round(ev['compression'] * 100)}% of the source length, {round(ev['verbatim'] * 100)}% verbatim")
    return EXAMPLE_PROMPT.format(metrics=metrics, used=", ".join(used) or "(none)", skipped=", ".join(skipped) or "(none)",
                                 pairs=pairs or "(no source documents)",
                                 unmatched="\n".join(f"- {u['text']}" for u in ev["unmatched"]) or "(none)",
                                 summary=citations.strip(ex["summary"]))


async def _notes(pid: str, ex: dict, ev: dict, provider: dict) -> tuple[str, bool]:
    """The LLM's analysis of one example, cached per example and model. Returns (notes, was_cached)."""
    f = _kdir(projects.get(pid)) / "examples" / ex["id"] / "notes.json"
    key = _key(ex, provider.get("model", ""))
    if f.exists():
        cached = json.loads(f.read_text("utf-8"))
        if cached.get("key") == key:
            return cached["notes"], True
    notes = await llm.complete([{"role": "user", "content": _example_prompt(ex, ev)}], provider)
    f.write_text(json.dumps({"key": key, "notes": notes}, ensure_ascii=False), "utf-8")
    return notes, False


async def learn(pid: str, provider: dict, progress=lambda msg: None) -> dict:
    """Draft a profile from all examples. Nothing is saved; the user confirms it (save_profile)."""
    project = projects.get(pid)
    examples = list_examples(pid)
    if not examples:
        raise ValueError("Add at least one example (source document + expected summary) first.")
    evs, analyses, new = [], [], 0
    for i, ex in enumerate(examples, 1):
        progress(f"Measuring example {i} of {len(examples)}: {ex['name']}…")
        ev = evidence(pid, ex)
        evs.append(ev)
        progress(f"Analyzing example {i} of {len(examples)}: {ex['name']}…")
        notes, cached = await _notes(pid, ex, ev, provider)
        new += 0 if cached else 1
        analyses.append(f"### Example {i}: {ex['name']}\n{notes}")
    progress("Combining everything into a profile…")
    m = aggregate(evs)
    measured = measured_markdown(m)
    prompt = MERGE_PROMPT.format(n=len(examples), measured=measured,
                                 instructions=project["instructions"].strip() or "(none)",
                                 analyses="\n\n".join(analyses))
    out = await llm.complete([{"role": "user", "content": prompt}], provider)
    learned, suggestions = _split_suggestions(out)
    return {"learned": learned, "suggestions": suggestions, "measured_md": measured, "targets": m["targets"],
            "pairs": m["pairs"], "examples": len(examples), "newly_analyzed": new}


def _split_suggestions(text: str) -> tuple[str, str]:
    m = re.search(r"^#{1,3}\s*Suggested additions.*$", text, re.M | re.I)
    if not m:
        return text.strip(), ""
    return text[:m.start()].strip(), text[m.end():].strip()


# ---------------------------------------------------------------- prompt

def prompt_section(project: dict, max_example_chars: int = 2500) -> str:
    """What a chat in this project gets from the knowledge base: the profile, the source map, and a few
    rewrite examples (source passage -> summary sentence), or one example summary for profiles without them."""
    prof = get_profile(project)
    if not prof:
        return ""
    out = ("\n\nThis project has a confirmed summary profile. When the user asks for a summary, follow it: the same "
           "structure, length, coverage of source sections, share of rewriting, tone and formatting. Use 'draws on' to "
           "find the material for each section. If information listed under 'Required information' is missing from "
           "the documents, ask the user for it before drafting (list exactly what is missing); if the user asks you to "
           "go ahead anyway, mark the gaps as [MISSING: ...]. Pinned rules override everything else in the profile."
           f"\n<summary_profile>\n{prof['text']}\n</summary_profile>")
    if prof.get("pairs"):
        ex = "\n\n".join(f"Source ({p['source_section']}): {p['source_sentence'][:400]}\n"
                          f"Summary ({p['section'] or 'summary'}): {citations.strip(p['text'])}" for p in prof["pairs"])
        out += ("\nHow the analysts rewrite source text into summary text (examples from other documents: copy the "
                f"style, not the facts):\n<rewrite_examples>\n{ex}\n</rewrite_examples>")
        return out
    try:
        examples = list_examples(project["id"])
    except FileNotFoundError:
        examples = []
    if examples:
        ex = citations.strip(examples[-1]["summary"])
        if len(ex) > max_example_chars:
            ex = ex[:max_example_chars] + "\n[...]"
        out += ("\nAn approved summary from this project, as a format reference only (do not reuse its facts; its "
                "source citations were removed, so still cite every fact as instructed):"
                f"\n<example_summary>\n{ex}\n</example_summary>")
    return out


# ---------------------------------------------------------------- checking a draft (C)

def check_draft(project: dict, draft_md: str, docs: list[tuple[str, dict]]) -> list[dict]:
    """Compare a drafted summary with the profile's measured targets. Each check: {label, ok, text}
    where ok is True (fits), False (outside the target / problem) or None (information)."""
    prof = get_profile(project)
    t = (prof or {}).get("targets") or {}
    if not t or len(alignment.summary_points(draft_md)) < 2:
        return []
    ev = alignment.align(draft_md, docs)
    out = []
    if t.get("words"):
        lo, hi = t["words"]
        out.append({"label": "Length", "ok": lo <= ev["words"] <= hi, "text": f"{ev['words']} words (target {lo}–{hi})"})
    if t.get("sections"):
        have = {_norm(s["name"]) for s in ev["sections"]}
        missing = [s for s in t["sections"] if _norm(s) not in have]
        out.append({"label": "Sections", "ok": not missing,
                    "text": "all expected sections" if not missing else "missing: " + ", ".join(missing)})
    if ev["numbers"]:
        bad = [n for r in ev["sentences"] for n in r["numbers_missing"]] + \
              [x["raw"] for u in ev["unmatched"] for x in alignment.numbers(u["text"])]
        out.append({"label": "Figures", "ok": not bad,
                    "text": f"all {ev['numbers']} found in the sources" if not bad else "not found in sources: " + ", ".join(bad[:6])})
    if t.get("verbatim") and ev["sentences"]:
        lo, hi = t["verbatim"]
        v = round(ev["verbatim"] * 100)
        out.append({"label": "Rewriting", "ok": lo <= v <= hi, "text": f"{v}% verbatim (examples {lo}–{hi}%)"})
    if t.get("required_sources"):
        present = {k for k, c in ev["coverage"].items() if c["words"] >= 15}
        unused = [s for s in t["required_sources"] if s in present and not ev["coverage"][s]["used"]]
        # With a single example, "always used" is weak evidence: report it, don't flag it.
        out.append({"label": "Coverage", "ok": (not unused) if t.get("examples", 1) >= 2 else (True if not unused else None),
                    "text": "uses every always-used section" if not unused else "not used: " + ", ".join(unused)})
    if ev["unmatched"] and len(ev["unmatched"]) > t.get("unsourced", 0) + 1:
        out.append({"label": "Unsourced", "ok": None,
                    "text": f"{len(ev['unmatched'])} sentence(s) with no clear source"})
    return out
