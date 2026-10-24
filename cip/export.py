"""Word summary per company, one Excel review checklist per run, optional highlighted PDF."""
import re

import pymupdf
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config

PART_TITLES = {"situation_recap": "Situation Recap", "recent_development": "Recent Development"}


def _add_rich(paragraph, text, highlight=False):
    """Write text with **bold** markdown into a docx paragraph."""
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        run.bold = i % 2 == 1
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def write_docx(company, summary, parsed, path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(f"{company} — Monthly Change in Provision", level=1)

    by_id = parsed.by_id()
    cited = []
    for part, title in PART_TITLES.items():
        doc.add_heading(title, level=2)
        for p in summary.get(part, []):
            para = doc.add_paragraph(style="List Bullet")
            _add_rich(para, p["text"], highlight=config.HIGHLIGHT_FLAGS_IN_WORD and p.get("flagged"))
            ref = para.add_run(f" [{', '.join(p['cites'])}]")
            ref.font.superscript = True
            ref.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            cited += [c for c in p["cites"] if c not in cited]

    doc.add_heading("Sources", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for cell, h in zip(t.rows[0].cells, ("ID", "Location in memo", "Paragraph text")):
        cell.text = h
    for c in sorted(cited, key=lambda x: [p.id for p in parsed.paragraphs].index(x) if x in by_id else 999):
        row = t.add_row().cells
        para = by_id.get(c)
        row[0].text = c
        row[1].text = para.location if para else "?"
        row[2].text = para.text if para else "(unknown ID)"
        for cell in row:
            for r in cell.paragraphs[0].runs:
                r.font.size = Pt(8.5)
    doc.save(path)


def write_highlighted_pdf(summary, parsed, path):
    """Copy of the source PDF with every cited paragraph highlighted and labelled with its ID."""
    doc = pymupdf.open(parsed.path)
    by_id = parsed.by_id()
    cited = {c for part in PART_TITLES for p in summary.get(part, []) for c in p["cites"]}
    for cid in cited:
        para = by_id.get(cid)
        if not para:
            continue
        for k, (pno, rect) in enumerate(para.rects):
            page = doc[pno]
            annot = page.add_highlight_annot(pymupdf.Rect(rect))
            annot.set_info(content=cid)
            annot.update()
            if k == 0:
                page.insert_text((max(rect[0] - 46, 8), rect[1] + 8), cid, fontsize=7, color=(0.8, 0, 0))
    doc.save(path)


def write_review_xlsx(results, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    head = ["Company", "Point", "Summary text", "Cites", "Cited location(s)",
            "Numbers not found in source", "LLM check", "LLM issue", "Reviewer OK (Y/N)", "Reviewer comment"]
    ws.append(head)
    red = PatternFill("solid", fgColor="F8D7DA")
    for r in results:
        if r.get("error"):
            ws.append([r["company"], "ERROR", r["error"]])
            continue
        by_id = r["parsed"].by_id()
        for part in PART_TITLES:
            for p in r["summary"].get(part, []):
                ch = p.get("checks", {})
                ws.append([
                    r["company"], p["id"], p["text"].replace("**", ""), ", ".join(p["cites"]),
                    "\n".join(by_id[c].location for c in p["cites"] if c in by_id),
                    ", ".join(ch.get("numbers_not_found", [])), ch.get("llm_verdict", ""),
                    ch.get("llm_issue", ""), "", "",
                ])
                if p.get("flagged"):
                    for cell in ws[ws.max_row]:
                        cell.fill = red
    widths = [22, 7, 70, 14, 34, 20, 18, 40, 12, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(path)
