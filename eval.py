"""
Step 3 of prompt tuning: score the AI against historical human summaries.

    eval/<case>/memo.pdf
    eval/<case>/human_summary.md
    python eval.py [--note "what I changed"]   ->  eval/reports/<timestamp>.xlsx  (+ a line in eval/reports/history.csv)

Cases whose folder name is also in prompts/examples/ are skipped (no testing on the examples).
"""
import argparse
import csv
import json
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import config
from llm_client import call_llm, parse_json
from parse import parse_pdf
from summarize import summarize
from verify import verify

DIMS = ("coverage", "accuracy", "style")


def as_markdown(summary):
    out = []
    for part, title in (("situation_recap", "Situation Recap"), ("recent_development", "Recent Development")):
        out.append(f"**{title}**\n")
        out += [f"- {p['text']}" for p in summary.get(part, [])]
        out.append("")
    return "\n".join(out)


def run_case(case):
    t0 = time.time()
    try:
        parsed = parse_pdf(case / "memo.pdf")
        summary = verify(summarize(parsed), parsed)
        ai_md = as_markdown(summary)
        human = (case / "human_summary.md").read_text(encoding="utf-8")
        judge_sys = (config.PROMPT_DIR / "meta" / "judge.md").read_text(encoding="utf-8")
        judge_user = (f"<memo>\n{parsed.to_prompt_text()}\n</memo>\n\n<reference_summary>\n{human}\n"
                      f"</reference_summary>\n\n<ai_summary>\n{ai_md}\n</ai_summary>")
        score = parse_json(call_llm(judge_sys, judge_user, task="judge"))
        flagged = sum(p.get("flagged", False) for part in ("situation_recap", "recent_development")
                      for p in summary.get(part, []))
        print(f"  {case.name}: " + " ".join(f"{d}={score.get(d)}" for d in DIMS) + f"  ({time.time()-t0:.0f}s)")
        return {"case": case.name, "ai": ai_md, "human": human, "flagged": flagged, **score}
    except Exception as e:
        print(f"  {case.name}: ERROR {e}")
        return {"case": case.name, "error": str(e)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--note", default="", help="what changed in the prompt since the last run")
    args = ap.parse_args()

    example_names = {p.name for p in config.EXAMPLES_DIR.iterdir() if p.is_dir()} if config.EXAMPLES_DIR.exists() else set()
    cases = [c for c in sorted(config.EVAL_DIR.iterdir())
             if c.is_dir() and (c / "memo.pdf").exists() and (c / "human_summary.md").exists()
             and c.name not in example_names]
    if not cases:
        raise SystemExit("No eval cases. Put eval/<case>/memo.pdf + human_summary.md")

    print(f"Evaluating {len(cases)} case(s) [LLM_MODE={config.LLM_MODE}]")
    with ThreadPoolExecutor(max_workers=config.MAX_WORKERS) as pool:
        rows = list(pool.map(run_case, cases))

    ok = [r for r in rows if "error" not in r]
    avg = {d: round(sum(r.get(d, 0) for r in ok) / len(ok), 2) if ok else 0 for d in DIMS}
    print("Average: " + "  ".join(f"{d}={v}" for d, v in avg.items()))

    reports = config.EVAL_DIR / "reports"
    reports.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")

    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"
    head = ["Case", *DIMS, "Flagged points", "Missing key points", "Errors", "Style gaps", "AI summary", "Human summary"]
    ws.append(head)
    for r in rows:
        if "error" in r:
            ws.append([r["case"], "ERROR", r["error"]])
            continue
        ws.append([r["case"], *(r.get(d) for d in DIMS), r["flagged"],
                   "\n".join(r.get("missing_key_points", [])), "\n".join(r.get("errors", [])),
                   "\n".join(r.get("style_gaps", [])), r["ai"], r["human"]])
    ws.append(["AVERAGE", *avg.values()])
    for i, w in enumerate([18, 9, 9, 9, 9, 40, 40, 40, 70, 70]):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(reports / f"{stamp}.xlsx")

    hist = reports / "history.csv"
    new = not hist.exists()
    with hist.open("a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["timestamp", "n_cases", *DIMS, "note"])
        w.writerow([stamp, len(ok), *avg.values(), args.note])
    print(f"Report: {reports / (stamp + '.xlsx')}\nHistory: {hist}")


if __name__ == "__main__":
    main()
