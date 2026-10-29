"""
Files the model creates in its answers, and writing them to disk.

The model writes a file as a block in its answer:
    <file path="summaries/harborline.docx">
    ...content (Markdown for .md, .docx, .xlsx and .pptx)...
    </file>
`split_answer` separates those blocks from the prose; `to_bytes` turns a block into file bytes:
.docx with python-docx, .xlsx with openpyxl (a sheet per ## heading, tables as cells), .pptx with python-pptx
(a slide per ## heading). Everything is converted locally; nothing leaves the computer.
"""
import io
import re
from pathlib import Path

FILE_TYPES = {".md", ".txt", ".csv", ".json", ".docx", ".xlsx", ".pptx"}
MARKDOWN_TYPES = {".md", ".txt", ".docx", ".xlsx", ".pptx"}    # written in Markdown, may contain citations

FILE_INSTRUCTIONS = """
Creating files: you CAN create files for the user, but only when the user asks for a file (e.g. "save it as Word", "give me a download", "export to CSV"). Otherwise answer in the chat as normal text with citations. To create a file, write it inside your answer as
<file path="NAME.EXT">
full file content
</file>
{where}
- Supported types: .md, .txt, .csv, .json, .docx, .xlsx and .pptx. For .docx, .xlsx and .pptx write the content in Markdown; it is converted:
  - .docx (Word): headings, bullet lists, **bold**, tables.
  - .xlsx (Excel): each "## Heading" starts a new sheet named after it; put the data in Markdown tables (header row first). Write numbers plainly (1,234.5 / 12.5% / $3.2) without other text in the cell, and put citations in a last column named "Source", not inside number cells.
  - .pptx (PowerPoint): "# Title" makes the title slide (a line of text under it becomes the subtitle); each "## Heading" starts a slide with that title; use short bullet points (at most about 6 per slide, nested bullets allowed) or one Markdown table per slide. A line starting with "Notes:" becomes the slide's speaker notes.
- Use a short descriptive file name; subfolders are allowed (e.g. path="drafts/summary.docx").
- Put the complete content in the block. Outside the block, say in one sentence what the file contains.
- Keep the citations inside the file content. The app shows the file's text in the chat with clickable citations, followed by a download/preview box, so do not repeat the content outside the block.
- If the user asks for a file of something you already wrote (e.g. "make that a Word file"), put that same content in the file block.
- Never say that you cannot create files."""

WHERE_PROJECT = "The user saves it to the project folder with one click (or downloads it)."
WHERE_CHAT = "The user downloads it with one click."

_BLOCK = re.compile(r'<file\s+(?:path|name)\s*=\s*["\']([^"\']+)["\']\s*>\n?(.*?)\n?</file>', re.S | re.I)
_OPEN = re.compile(r'<file\s+(?:path|name)\s*=\s*["\']([^"\']+)["\']\s*>\n?', re.I)


def clean_path(path: str) -> str:
    """Relative path with a supported suffix (unknown suffixes become .md)."""
    p = Path(path.strip().replace("\\", "/").lstrip("/"))
    parts = [re.sub(r'[<>:"|?*]', "_", s) for s in p.parts if s not in ("", ".", "..")]
    p = Path(*parts) if parts else Path("output.md")
    return str(p if p.suffix.lower() in FILE_TYPES else p.with_suffix(".md"))


def split_answer(text: str) -> list[dict]:
    """[{"kind": "text", "text"} | {"kind": "file", "path", "content", "complete"}] in answer order."""
    parts, at = [], 0
    for m in _BLOCK.finditer(text):
        if m.start() > at:
            parts.append({"kind": "text", "text": text[at:m.start()]})
        parts.append({"kind": "file", "path": clean_path(m.group(1)), "content": m.group(2).strip("\n") + "\n",
                      "complete": True})
        at = m.end()
    rest = text[at:]
    m = _OPEN.search(rest)
    if m:                                   # a block still being written (streaming)
        if m.start():
            parts.append({"kind": "text", "text": rest[:m.start()]})
        parts.append({"kind": "file", "path": clean_path(m.group(1)), "content": rest[m.end():], "complete": False})
    elif rest:
        parts.append({"kind": "text", "text": rest})
    return parts


def to_bytes(path: str, content: str) -> bytes:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return markdown_to_docx(content)
    if suffix == ".xlsx":
        return markdown_to_xlsx(content)
    if suffix == ".pptx":
        return markdown_to_pptx(content)
    return content.encode("utf-8")


# ---------------------------------------------------------------- Markdown -> Word

def _add_runs(paragraph, text: str):
    """Add text to a python-docx paragraph, honouring **bold** and *italic*."""
    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def markdown_to_docx(md: str) -> bytes:
    import docx

    doc = docx.Document()
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?\s*:?-{2,}", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2).strip().strip("*"), level=min(len(m.group(1)), 4))
        elif re.match(r"^[-*+•]\s+", stripped):
            indent = len(line) - len(line.lstrip())
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            _add_runs(doc.add_paragraph(style=style), re.sub(r"^[-*+•]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
        elif re.fullmatch(r"-{3,}|\*{3,}", stripped):
            doc.add_paragraph()
        else:
            _add_runs(doc.add_paragraph(), stripped)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Markdown blocks (for Excel and PowerPoint)

_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}")


def _md_blocks(md: str):
    """Yield ("heading", level, text) | ("table", rows) | ("bullet", level, text) | ("text", text) | ("notes", text)."""
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and _TABLE_RULE.match(lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _TABLE_RULE.match(lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            yield "table", rows
            continue
        indent = len(line) - len(line.lstrip())
        if m := re.match(r"^(#{1,6})\s+(.*)$", stripped):
            yield "heading", len(m.group(1)), m.group(2).strip().strip("*").strip()
        elif re.match(r"^[-*+•]\s+", stripped):
            yield "bullet", min(indent // 2, 4), re.sub(r"^[-*+•]\s+", "", stripped)
        elif re.match(r"^\d+[.)]\s+", stripped):
            yield "bullet", min(indent // 2, 4), stripped
        elif m := re.match(r"^\**notes?\**\s*:\**\s*(.*)$", stripped, re.I):
            yield "notes", m.group(1)
        else:
            yield "text", stripped
        i += 1


def _plain(text: str) -> str:
    return re.sub(r"\*\*([^*]+)\*\*|\*([^*\s][^*]*)\*", lambda m: m.group(1) or m.group(2), text)


# ---------------------------------------------------------------- Markdown -> Excel

_NUMBER = re.compile(r"^(?P<neg>\(|-)?(?P<cur>\$)?(?P<num>\d{1,3}(?:,\d{3})+|\d+)(?P<dec>\.\d+)?(?P<pct>%)?\)?$")


def _cell_value(text: str):
    """A Markdown cell as an Excel value: plain numbers, amounts and percentages become numbers."""
    s = _plain(text).strip()
    m = _NUMBER.match(s.replace(" ", ""))
    if not m or (m.group("neg") == "(" and not s.endswith(")")):
        return s, None
    value = float(m.group("num").replace(",", "") + (m.group("dec") or ""))
    value = -value if m.group("neg") else value
    decimals = len(m.group("dec")) - 1 if m.group("dec") else 0
    zeros = ("." + "0" * decimals) if decimals else ""
    if m.group("pct"):
        return value / 100, f"0{zeros}%"
    if m.group("cur"):
        return value, f"$#,##0{zeros}"
    if "," in m.group("num"):
        return value, f"#,##0{zeros}"
    return (int(value) if not decimals else value), None


def markdown_to_xlsx(md: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    used: set[str] = set()
    state = {"row": 1, "fresh": True, "intro": False, "widths": {}}
    bold, header_fill = Font(bold=True), PatternFill("solid", fgColor="DDE5F4")

    def sheet_name(text: str) -> str:
        base = re.sub(r"[\[\]:*?/\\]", " ", _plain(text)).strip()[:31] or "Sheet"
        name, n = base, 2
        while name.lower() in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name.lower())
        return name

    def new_sheet(title: str):
        nonlocal ws
        if state["fresh"] or state["intro"]:     # nothing (or only the title) written yet: name the current sheet
            ws.title = sheet_name(title)
            state["intro"] = False
            if not state["fresh"]:
                return
        else:
            finish(ws)
            ws = wb.create_sheet(sheet_name(title))
        state.update(row=1, fresh=True, widths={})

    def put(col: int, text: str, font=None, fill=None, indent=0):
        value, fmt = _cell_value(text)
        cell = ws.cell(row=state["row"], column=col, value=value)
        if fmt:
            cell.number_format = fmt
        if font or text.strip().startswith("**") and text.strip().endswith("**"):
            cell.font = font or bold
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top", indent=indent)
        width = max(len(line) for line in str(value).split("\n")) if value != "" else 0
        state["widths"][col] = max(state["widths"].get(col, 8), min(width + 2, 60))
        state["fresh"] = False

    def finish(sheet):
        for col, width in state["widths"].items():
            sheet.column_dimensions[get_column_letter(col)].width = width

    title_done = False
    for block in _md_blocks(md):
        kind = block[0]
        if kind == "heading":
            level, text = block[1], block[2]
            if level == 1 and not title_done and state["fresh"]:
                put(1, text, font=Font(bold=True, size=14))
                state["row"] += 1
                state["intro"] = title_done = True
            elif level <= 2:
                intro = state["intro"]
                new_sheet(text)
                if intro:
                    state["row"] += 1
            else:
                if not state["fresh"]:
                    state["row"] += 1
                put(1, text, font=bold)
                state["row"] += 1
        elif kind == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            for r, row in enumerate(rows):
                for c in range(width):
                    put(c + 1, row[c] if c < len(row) else "", font=bold if r == 0 else None,
                        fill=header_fill if r == 0 else None)
                state["row"] += 1
            if state["row"] - len(rows) == 1 and ws.freeze_panes is None:
                ws.freeze_panes = "A2"           # the table starts the sheet: keep its header visible
            state["row"] += 1
        elif kind == "bullet" and ws.title.lower().startswith("sources") and (m := re.match(r"^(\d+)[.)]\s+(.*)", block[2])):
            put(1, m.group(1))
            put(2, m.group(2))
            state["row"] += 1
        elif kind == "bullet":
            put(1, "• " + block[2], indent=block[1] * 2)
            state["row"] += 1
        else:                                    # text and notes
            put(1, block[-1])
            state["row"] += 1
            continue
        if kind != "heading" or block[1] > 2:
            state["intro"] = False               # the sheet has its own content now
    finish(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Markdown -> PowerPoint

_SLIDE_BUDGET = 12          # rough lines of text that fit on a slide at the default font sizes


def _pptx_runs(paragraph, text: str, size=None):
    from pptx.util import Pt

    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
        if not piece:
            continue
        run = paragraph.add_run()
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            run.text, run.font.bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run.text, run.font.italic = piece[1:-1], True
        else:
            run.text = piece
        if size:
            run.font.size = Pt(size)


def markdown_to_pptx(md: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    TITLE, CONTENT, TITLE_ONLY = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]
    margin, top = Inches(0.6), Inches(1.5)
    state = {"slide": None, "frame": None, "used": 0, "title": "", "title_slide": None, "sources": False}

    def place(shape, left, top_, width, height):
        # The default template is 4:3; placeholders are moved explicitly to fit the 16:9 slide.
        shape.left, shape.top, shape.width, shape.height = left, top_, width, height

    def fit_title(slide, text: str):
        slide.shapes.title.text = _plain(text)
        place(slide.shapes.title, margin, Inches(0.35), prs.slide_width - 2 * margin, Inches(1.0))
        for p in slide.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(30 if len(text) < 60 else 24)

    def start(title: str, layout=CONTENT):
        slide = prs.slides.add_slide(layout)
        fit_title(slide, title)
        frame = None
        if layout is CONTENT:
            body = slide.placeholders[1]
            place(body, margin, top, prs.slide_width - 2 * margin, prs.slide_height - top - Inches(0.5))
            frame = body.text_frame
            frame.word_wrap = True
        state.update(slide=slide, frame=frame, used=0, title=title)
        return slide

    def text_frame(cost: int):
        """The body to write into, continuing on a new slide when this one is full."""
        if state["frame"] is None or state["used"] + cost > _SLIDE_BUDGET:
            if state["slide"] is None or state["slide"] is state["title_slide"]:
                start("Overview" if state["title_slide"] else "Summary")
            else:
                start(re.sub(r"( \(cont\.\))?$", " (cont.)", state["title"], count=1))
        state["used"] += cost
        return state["frame"]

    def paragraph(text: str, level: int = 0, bold: bool = False):
        small = state["sources"]
        cost = 1 + len(_plain(text)) // (150 if small else 95)
        frame = text_frame(cost)
        p = frame.paragraphs[0] if len(frame.paragraphs) == 1 and not frame.paragraphs[0].runs else frame.add_paragraph()
        p.level = level
        _pptx_runs(p, f"**{text}**" if bold else text, size=12 if small else (20 if level == 0 else 17))

    def table(rows: list[list[str]]):
        if state["slide"] is None or state["used"] > 0:
            start((state["title"] or "Summary") + (" (cont.)" if state["used"] else ""), TITLE_ONLY)
        elif state["frame"] is not None:         # an empty content slide: use its space for the table
            state["slide"].shapes._spTree.remove(state["slide"].placeholders[1]._element)
        slide, width = state["slide"], max(len(r) for r in rows)
        height = min(Inches(0.4) * len(rows), prs.slide_height - top - Inches(0.4))
        shape = slide.shapes.add_table(len(rows), width, margin, top, prs.slide_width - 2 * margin, Emu(int(height)))
        size = 14 if len(rows) <= 8 else 11
        for r, row in enumerate(rows):
            for c in range(width):
                cell = shape.table.cell(r, c)
                cell.text = ""
                _pptx_runs(cell.text_frame.paragraphs[0], row[c] if c < len(row) else "", size=size)
        state.update(frame=None, used=_SLIDE_BUDGET)

    for block in _md_blocks(md):
        kind = block[0]
        if kind == "heading":
            level, text = block[1], block[2]
            if level == 1 and state["title_slide"] is None and state["slide"] is None:
                slide = prs.slides.add_slide(TITLE)
                slide.shapes.title.text = _plain(text)
                place(slide.shapes.title, margin, Inches(2.2), prs.slide_width - 2 * margin, Inches(1.6))
                place(slide.placeholders[1], margin, Inches(4.0), prs.slide_width - 2 * margin, Inches(1.4))
                state["title_slide"] = slide
                state.update(slide=slide, frame=None, used=_SLIDE_BUDGET, title=text)
            elif level <= 2:
                state["sources"] = text.strip().lower() in ("sources", "references")
                start(text)
            else:
                paragraph(text, bold=True)
        elif kind == "text" and state["slide"] is not None and state["slide"] is state["title_slide"] \
                and len(state["title_slide"].placeholders) > 1 and not state["title_slide"].placeholders[1].text:
            state["title_slide"].placeholders[1].text = _plain(block[1])     # subtitle
        elif kind == "notes":
            if state["slide"] is not None:
                notes = state["slide"].notes_slide.notes_text_frame
                notes.text = (notes.text + "\n" if notes.text else "") + _plain(block[1])
        elif kind == "table":
            table(block[1])
        elif kind == "bullet":
            paragraph(block[2], level=block[1])
        else:
            paragraph(block[1])
    if not len(prs.slides):
        start("Summary")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
